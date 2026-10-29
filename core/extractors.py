
from pathlib import Path


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
}


class TextExtractor:
    """استخراج متن از فایل‌های پشتیبانی‌شده، بدون تولید محتوای حقوقی."""

    def extract(self, file_path):
        path = Path(file_path)

        if not path.is_file():
            raise FileNotFoundError(
                f"فایل پیدا نشد: {path}"
            )

        extension = path.suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"نوع فایل پشتیبانی نمی‌شود: {extension}"
            )

        if extension == ".pdf":
            return self._extract_pdf(path)

        if extension == ".docx":
            return self._extract_docx(path)

        if extension in {
            ".png", ".jpg", ".jpeg",
            ".tif", ".tiff", ".bmp"
        }:
            return self._extract_image(path)

        raise ValueError("استخراج متن از این فایل امکان‌پذیر نیست.")

    @staticmethod
    def _extract_pdf(path):
        import pdfplumber

        pages = []

        with pdfplumber.open(str(path)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""

                if text.strip():
                    pages.append(
                        f"[صفحه {page_number}]\n{text.strip()}"
                    )

        result = "\n\n".join(pages).strip()

        if not result:
            raise ValueError(
                "متن قابل استخراج در PDF پیدا نشد؛ "
                "ممکن است فایل اسکن‌شده باشد و به OCR نیاز داشته باشد."
            )

        return result

    @staticmethod
    def _extract_docx(path):
        from docx import Document

        document = Document(str(path))
        parts = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]

                if any(cells):
                    parts.append(" | ".join(cells))

        result = "\n".join(parts).strip()

        if not result:
            raise ValueError(
                "متنی در فایل Word پیدا نشد."
            )

        return result

    @staticmethod
    def _extract_image(path):
        from PIL import Image
        import pytesseract

        with Image.open(str(path)) as image:
            text = pytesseract.image_to_string(
                image,
                lang="fas+eng"
            )

        result = text.strip()

        if not result:
            raise ValueError(
                "OCR متنی استخراج نکرد. "
                "نصب موتور Tesseract و زبان فارسی را بررسی کنید."
            )

        return result
